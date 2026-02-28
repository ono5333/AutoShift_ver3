"""
AutoShift - シフト表示API
生成されたシフトの表示とダウンロード機能
2026年2月9日
"""

from flask import Blueprint, jsonify, request, render_template, send_file
from pathlib import Path
import json
import csv
import io
import sys
from datetime import datetime
from typing import Dict, Any

# プロジェクトパスを追加
sys.path.insert(0, str(Path(__file__).parent.parent))

from optimizer.solver import ShiftOptimizer
from utils.shift_display import ShiftDisplayManager
from utils.yaml_handler import (
    load_project_carryover_for_month
)

shift_result_bp = Blueprint('shift_result', __name__)


def _previous_month_str(month: str) -> str:
    dt = datetime.strptime(month, "%Y-%m")
    prev_year = dt.year if dt.month > 1 else dt.year - 1
    prev_month = dt.month - 1 if dt.month > 1 else 12
    return f"{prev_year:04d}-{prev_month:02d}"


SHIFT_LABELS = {'日', '夜', '明', '休', '希', '有'}


def _normalize_prev_month_last_shifts(raw: Any) -> Dict[int, str]:
    if not isinstance(raw, dict):
        return {}
    normalized: Dict[int, str] = {}
    for k, v in raw.items():
        try:
            staff_id = int(k)
        except (ValueError, TypeError):
            continue
        if isinstance(v, str) and v in SHIFT_LABELS:
            normalized[staff_id] = v
    return normalized


def _expected_history_csv_path(target_month: str) -> str:
    prev = _previous_month_str(target_month).replace('-', '_')
    return str(Path("data") / "history" / f"shift_history_{prev}.csv")


def _ensure_carryover_for_month(target_month: str) -> Dict[str, Any]:
    entry = load_project_carryover_for_month(target_month)
    if not isinstance(entry, dict):
        return {"ready": False, "source": "none", "night_staff_ids": [], "prev_month_last_shifts": {}}

    last_shifts = _normalize_prev_month_last_shifts(entry.get('prev_month_last_shifts', {}))
    if last_shifts:
        night_staff_ids = sorted([sid for sid, s in last_shifts.items() if s == '夜'])
        return {
            "ready": True,
            "source": entry.get('source', 'history_csv'),
            "night_staff_ids": night_staff_ids,
            "prev_month_last_shifts": {str(k): v for k, v in sorted(last_shifts.items())}
        }

    return {"ready": False, "source": "none", "night_staff_ids": [], "prev_month_last_shifts": {}}


# ====================

# API エンドポイント
# ====================

@shift_result_bp.route('/api/shift/display/generate/<month>', methods=['POST'])
def generate_and_display_shift(month):
    """
    シフト生成して表示API
    
    Args:
        month: 対象月 (YYYY-MM形式)
        
    Returns:
        JSON: 生成されたシフト情報
    """
    try:
        # 月の形式確認
        try:
            datetime.strptime(month, '%Y-%m')
        except ValueError:
            return jsonify({
                'status': 'error',
                'message': '月の形式が正しくありません (YYYY-MM形式で指定してください)'
            }), 400

        carryover_state = _ensure_carryover_for_month(month)
        if not carryover_state.get("ready"):
            return jsonify({
                'status': 'need_carryover',
                'message': (
                    f'前月履歴CSVが不足しています。'
                    f' {_expected_history_csv_path(month)} を用意してください。'
                ),
                'data': {
                    'month': month,
                    'previous_month': _previous_month_str(month),
                    'expected_history_csv': _expected_history_csv_path(month)
                }
            }), 409

        # 最適化実行
        # print(f"[RUN] {month} のシフト最適化を開始...")
        optimizer = ShiftOptimizer(
            month,
            carryover_override={
                'night_staff_ids': carryover_state.get('night_staff_ids', []),
                'prev_month_last_shifts': carryover_state.get('prev_month_last_shifts', {})
            }
        )
        result = optimizer.optimize()

        # 実行可能解がない場合は、表示データを作らずエラー返却
        if result.solver_status in ('INFEASIBLE', 'UNKNOWN'):
            return jsonify({
                'status': 'error',
                'message': (
                    f'{month}はシフトを作成できませんでした '
                    f'(solver_status={result.solver_status})。'
                    'ルールまたは希望休を見直してください。'
                ),
                'data': {
                    'month': month,
                    'solver_status': result.solver_status,
                    'solver_status_code': result.solver_status_code,
                    'solver_status_name': result.solver_status_name,
                    'solver_time': round(result.solver_time, 4),
                    'diagnosis': result.diagnosis
                }
            }), 422
        
        # 表示マネージャー初期化
        display = ShiftDisplayManager()
        
        # シフト表データ生成
        shift_table = display.generate_shift_table_data(result, optimizer.staff_list, optimizer.dates)
        
        # 統計データ生成
        statistics = display.generate_statistics_data(result, optimizer.staff_list, optimizer.dates)
        
        # 違反情報
        violations = []
        if result.violations:
            staff_name_map = {s.id: s.name for s in optimizer.staff_list}
            for violation in result.violations:
                staff_id = violation.get('staff_id')
                if not isinstance(staff_id, int):
                    staff_id = None
                violations.append({
                    'type': violation.get('type', 'unknown'),
                    'rank': violation.get('rank') or (str(violation.get('type', '')).replace('Rank', '') if violation.get('type') else None),
                    'rule_id': violation.get('rule_id') or violation.get('constraint_name'),
                    'constraint_name': violation.get('constraint_name'),
                    'message': violation.get('message', ''),
                    'detail': violation.get('detail') or violation.get('message', ''),
                    'staff_id': violation.get('staff_id'),
                    'staff_name': violation.get('staff_name') or (staff_name_map.get(staff_id) if staff_id is not None else None),
                    'date': violation.get('date')
                })
        
        response_data = {
            'status': 'success',
            'message': f'{month}のシフトを生成しました',
            'data': {
                'month': month,
                'solver_status': result.solver_status,
                'solver_status_code': result.solver_status_code,
                'solver_status_name': result.solver_status_name,
                'solver_time': round(result.solver_time, 4),
                'shift_table': shift_table,
                'statistics': statistics,
                'violations': violations,
                'generated_at': datetime.now().isoformat()
            }
        }
        
        # print(f"[OK] シフト生成・表示データ作成完了: {result.solver_status} ({result.solver_time:.3f}秒)")
        return jsonify(response_data), 200
        
    except Exception as e:
        # print(f"[ERROR] シフト生成・表示エラー: {e}")
        return jsonify({
            'status': 'error',
            'message': f'シフト生成に失敗しました: {str(e)}'
        }), 500


@shift_result_bp.route('/api/shift/carryover/<month>/status', methods=['GET'])
def get_carryover_status(month):
    """前月履歴CSV由来の引継ぎ状況を返す。"""
    try:
        datetime.strptime(month, '%Y-%m')
    except ValueError:
        return jsonify({
            'status': 'error',
            'message': '月の形式が正しくありません (YYYY-MM形式で指定してください)'
        }), 400

    try:
        carryover_state = _ensure_carryover_for_month(month)
        if carryover_state.get("ready"):
            return jsonify({
                'status': 'success',
                'data': {
                    'month': month,
                    'previous_month': _previous_month_str(month),
                    'source': carryover_state.get('source', 'history_csv'),
                    'prev_month_last_shifts': carryover_state.get('prev_month_last_shifts', {})
                }
            }), 200

        return jsonify({
            'status': 'error',
            'message': (
                f'前月履歴CSVが不足しています。'
                f' {_expected_history_csv_path(month)} を用意してください。'
            ),
            'data': {
                'month': month,
                'previous_month': _previous_month_str(month),
                'expected_history_csv': _expected_history_csv_path(month)
            }
        }), 404
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'引継ぎ状況取得に失敗しました: {str(e)}'
        }), 500


@shift_result_bp.route('/api/shift/carryover/<month>', methods=['POST'])
def save_carryover_status(month):
    """CSV運用への切替により、手入力保存APIは非推奨。"""
    return jsonify({
        'status': 'error',
        'message': 'carryover手入力は廃止しました。前月履歴CSVを data/history に配置してください。'
    }), 410

@shift_result_bp.route('/api/shift/display/table/<month>', methods=['GET'])
def get_shift_table_data(month):
    """
    シフト表データ取得API（表示専用、生成は行わない）
    
    Args:
        month: 対象月 (YYYY-MM形式)
        
    Returns:
        JSON: 保存済みシフト表データまたはデータなしメッセージ
    """
    try:
        # 月の形式確認
        try:
            datetime.strptime(month, '%Y-%m')
        except ValueError:
            return jsonify({
                'status': 'error',
                'message': '月の形式が正しくありません (YYYY-MM形式で指定してください)'
            }), 400
        
        # 保存済み結果確認
        results_dir = Path(__file__).parent.parent / "results"
        result_file = results_dir / f"shift_result_{month.replace('-', '_')}.json"
        
        if result_file.exists():
            # 保存済みデータを読み込み
            with open(result_file, 'r', encoding='utf-8') as f:
                saved_data = json.load(f)
                
            # 表示マネージャー初期化
            display = ShiftDisplayManager()
            
            # 必要に応じてデータ形式を変換
            response_data = {
                'status': 'success',
                'message': f'{month}の保存済みシフトデータを読み込みました',
                'data': saved_data
            }
            
            return jsonify(response_data), 200
        else:
            # 保存済みデータが存在しない場合
            return jsonify({
                'status': 'no_data',
                'message': f'{month}のシフトデータがありません。生成ボタンを押してシフトを作成してください。'
            }), 200
            
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'シフト表データ取得に失敗しました: {str(e)}'
        }), 500

@shift_result_bp.route('/api/shift/display/export/<month>/<format>', methods=['GET'])
def export_shift_data(month, format):
    """
    シフトデータエクスポートAPI
    
    Args:
        month: 対象月 (YYYY-MM形式)
        format: エクスポート形式 (csv, json, excel)
        
    Returns:
        File: エクスポートファイル
    """
    try:
        # 最適化実行
        optimizer = ShiftOptimizer(month)
        result = optimizer.optimize()
        
        display = ShiftDisplayManager()
        
        if format.lower() == 'csv':
            # CSV生成
            csv_path = display.export_to_csv(result, optimizer.staff_list, optimizer.dates)
            return send_file(csv_path, as_attachment=True, 
                           download_name=f'shift_{month}.csv')
        
        elif format.lower() == 'json':
            # JSON生成
            shift_data = {
                'month': month,
                'generated_at': datetime.now().isoformat(),
                'solver_status': result.solver_status,
                'solver_time': result.solver_time,
                'shifts': {}
            }
            
            for (staff_id, date_obj), shift_type in result.shifts.items():
                date_str = date_obj.strftime('%Y-%m-%d')
                if date_str not in shift_data['shifts']:
                    shift_data['shifts'][date_str] = {}
                shift_data['shifts'][date_str][str(staff_id)] = shift_type.value
            
            # メモリ上でJSON作成
            json_buffer = io.StringIO()
            json.dump(shift_data, json_buffer, ensure_ascii=False, indent=2)
            json_buffer.seek(0)
            
            return send_file(
                io.BytesIO(json_buffer.getvalue().encode('utf-8')),
                mimetype='application/json',
                as_attachment=True,
                download_name=f'shift_{month}.json'
            )
        
        elif format.lower() == 'excel':
            # Excel生成（CSVと同じレイアウト + 曜日列の色分け）
            try:
                import openpyxl
                from openpyxl.utils import get_column_letter
                from openpyxl.styles import PatternFill, Alignment, Font
                import re

                # まずCSVと同じ表データを作る
                csv_path = display.export_to_csv(result, optimizer.staff_list, optimizer.dates)
                with open(csv_path, 'r', encoding='utf-8-sig', newline='') as f:
                    csv_rows = list(csv.reader(f))

                wb = openpyxl.Workbook()
                ws = wb.active
                if ws is None:
                    raise ValueError('Excelワークシートの初期化に失敗しました')
                ws.title = 'シフト表'

                for row in csv_rows:
                    ws.append(row)

                # 共通スタイル
                center = Alignment(horizontal='center', vertical='center')
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        cell.alignment = center

                header_fill = PatternFill(fill_type='solid', fgColor='F2F2F2')
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=1, column=c).font = Font(bold=True)
                    ws.cell(row=1, column=c).fill = header_fill

                # スタッフ行範囲（ヘッダー1行 + スタッフN行）
                staff_row_start = 2
                staff_row_end = 1 + len(optimizer.staff_list)

                # ID/スタッフ名/職種の背景色（フロントエンドに合わせる）
                default_staff_fill = PatternFill(fill_type='solid', fgColor='F8F9FA')
                junior_fill = PatternFill(fill_type='solid', fgColor='FFF9DB')  # 初級介護士
                bath_fill = PatternFill(fill_type='solid', fgColor='EAF7EA')    # お風呂
                for row_idx, staff in enumerate(optimizer.staff_list, start=staff_row_start):
                    if staff.staff_class == '初級介護士':
                        fill = junior_fill
                    elif staff.staff_class == 'お風呂':
                        fill = bath_fill
                    else:
                        fill = default_staff_fill
                    for col_idx in (1, 2, 3):  # A,B,C列
                        ws.cell(row=row_idx, column=col_idx).fill = fill

                # 列幅
                ws.column_dimensions['A'].width = 10
                ws.column_dimensions['B'].width = 16
                ws.column_dimensions['C'].width = 12
                for c in range(4, ws.max_column + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 9

                # 土日列を着色（ヘッダーが MM/DD(曜) でも MM/DD でも対応）
                sat_fill = PatternFill(fill_type='solid', fgColor='E6F4FF')
                sun_fill = PatternFill(fill_type='solid', fgColor='FFECEC')
                year = int(month.split('-')[0])
                for c in range(4, ws.max_column + 1):
                    raw = str(ws.cell(row=1, column=c).value or '').strip()
                    m = re.match(r'^(\d{2})/(\d{2})', raw)
                    if not m:
                        continue
                    mm = int(m.group(1))
                    dd = int(m.group(2))
                    try:
                        wd = datetime(year, mm, dd).weekday()
                    except Exception:
                        continue
                    fill = sat_fill if wd == 5 else (sun_fill if wd == 6 else None)
                    if fill:
                        # 合計系セルには土日色を付けない（スタッフ行まで）
                        for r in range(1, staff_row_end + 1):
                            ws.cell(row=r, column=c).fill = fill

                # 統計シート
                stats = display.generate_statistics_data(result, optimizer.staff_list, optimizer.dates)
                ws2 = wb.create_sheet('統計')
                ws2.append(['スタッフ名', '職種', '日勤', '夜勤', '明け', '公休', '有給', '希望休', '勤務日数', '勤務率(%)'])
                for s in stats.get('staff_stats', []):
                    ws2.append([
                        s.get('staff_name', ''),
                        s.get('staff_class', ''),
                        s.get('day_shifts', 0),
                        s.get('night_shifts', 0),
                        s.get('night_off_shifts', 0),
                        s.get('holiday_shifts', 0),
                        s.get('paid_holidays', 0),
                        s.get('request_holidays', 0),
                        s.get('work_days', 0),
                        s.get('work_ratio', 0),
                    ])
                for c in range(1, 11):
                    ws2.cell(row=1, column=c).font = Font(bold=True)
                    ws2.cell(row=1, column=c).fill = header_fill

                # Excel書き出し
                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                excel_buffer.seek(0)
                
                return send_file(
                    excel_buffer,
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    as_attachment=True,
                    download_name=f'shift_{month}.xlsx'
                )
                
            except ImportError as e:
                return jsonify({
                    'status': 'error',
                    'message': f'Excelエクスポートに必要なライブラリが不足しています: {e}'
                }), 500
        
        else:
            return jsonify({
                'status': 'error',
                'message': f'サポートされていないフォーマット: {format}'
            }), 400
            
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'エクスポートに失敗しました: {str(e)}'
        }), 500

@shift_result_bp.route('/api/shift/display/validate/<month>', methods=['GET'])
def validate_shift(month):
    """
    シフト妥当性検証API
    
    Args:
        month: 対象月 (YYYY-MM形式)
        
    Returns:
        JSON: 検証結果
    """
    try:
        # 最適化実行
        optimizer = ShiftOptimizer(month)
        result = optimizer.optimize()
        
        # 制約違反チェック
        violations = result.violations or []
        
        # 各種統計
        total_shifts = len(result.shifts)
        staff_count = len(optimizer.staff_list)
        date_count = len(optimizer.dates)
        
        # 勤務タイプ別集計
        shift_type_counts = {}
        for shift_type in result.shifts.values():
            type_name = shift_type.name
            shift_type_counts[type_name] = shift_type_counts.get(type_name, 0) + 1
        
        # 検証結果
        validation_result = {
            'is_valid': len(violations) == 0,
            'violation_count': len(violations),
            'violations': violations,
            'statistics': {
                'total_shifts': total_shifts,
                'staff_count': staff_count,
                'date_count': date_count,
                'shift_type_distribution': shift_type_counts
            },
            'performance': {
                'solver_status': result.solver_status,
                'solver_time': result.solver_time,
                'optimization_success': result.solver_status == 'OPTIMAL'
            }
        }
        
        return jsonify({
            'status': 'success',
            'data': validation_result
        }), 200
        
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'シフト検証に失敗しました: {str(e)}'
        }), 500

@shift_result_bp.route('/api/shift/display/compare/<month1>/<month2>', methods=['GET'])
def compare_shifts(month1, month2):
    """
    シフト比較API
    
    Args:
        month1: 比較対象月1 (YYYY-MM形式)
        month2: 比較対象月2 (YYYY-MM形式)
        
    Returns:
        JSON: 比較結果
    """
    try:
        # 両方の月のシフトを生成
        optimizer1 = ShiftOptimizer(month1)
        result1 = optimizer1.optimize()
        
        optimizer2 = ShiftOptimizer(month2)
        result2 = optimizer2.optimize()
        
        # 比較データ作成
        comparison = {
            'month1': {
                'month': month1,
                'solver_status': result1.solver_status,
                'solver_time': result1.solver_time,
                'violation_count': len(result1.violations or [])
            },
            'month2': {
                'month': month2,
                'solver_status': result2.solver_status,
                'solver_time': result2.solver_time,
                'violation_count': len(result2.violations or [])
            },
            'performance_comparison': {
                'faster_month': month1 if result1.solver_time < result2.solver_time else month2,
                'time_diff': abs(result1.solver_time - result2.solver_time),
                'both_optimal': (result1.solver_status == 'OPTIMAL' and result2.solver_status == 'OPTIMAL')
            }
        }
        
        return jsonify({
            'status': 'success',
            'data': comparison
        }), 200
        
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'シフト比較に失敗しました: {str(e)}'
        }), 500
